"""Read a version Excel export (possibly edited by hand) back into rows.

The export is written by the frontend (`exportQuotation.ts`); the header
names below are the contract between the two.
"""

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from ..models.quotation_line_item import QuotationItemStatus
from .dimension_parser import normalize_unit

MAX_ROWS = 2000
EXTRA_PREFIX = "extra:"
_SHEET = "Quotation"
_DEFAULT_UNIT = "lumpsum"

# Header text (lowercase) -> field on the line item.
_HEADERS = {
    "room": "room",
    "description": "description",
    "unit": "unit",
    "dimension": "raw_dimension",
    "qty": "qty",
    "rate": "rate",
    "initial price": "initial_price",
    "negotiated price": "negotiated_price",
    "est. time": "est_time",
    "status": "status",
    "remarks": "remarks",
}
_NUMBER_FIELDS = ("qty", "rate", "initial_price", "negotiated_price")
_MAX_LEN = {
    "room": 100,
    "description": 500,
    "unit": 20,
    "raw_dimension": 50,
    "est_time": 50,
    "remarks": 2000,
}


class ImportFileError(Exception):
    """The file cannot be read as a quotation; the message is user-facing."""


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _number(value: object, field: str, row: int) -> float | None:
    text = _text(value)
    if text is None:
        return None
    cleaned = text.replace("₹", "").replace(",", "").strip()
    try:
        number = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ImportFileError(
            f"Row {row}: '{text}' is not a number in {field.replace('_', ' ')}"
        ) from exc
    if number < 0:
        raise ImportFileError(
            f"Row {row}: {field.replace('_', ' ')} is negative"
        )
    return float(number)


def _status(value: object, row: int) -> QuotationItemStatus | None:
    text = _text(value)
    if text is None:
        return None
    try:
        return QuotationItemStatus(text.lower().replace(" ", "_"))
    except ValueError as exc:
        raise ImportFileError(f"Row {row}: unknown status '{text}'") from exc


def parse_version_xlsx(data: bytes) -> list[dict]:
    """Rows ready for `create_version_from_rows`."""
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises many types for bad files
        raise ImportFileError(
            "This is not a valid Excel (.xlsx) file"
        ) from exc

    sheet = (
        workbook[_SHEET] if _SHEET in workbook.sheetnames else workbook.active
    )
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise ImportFileError("The file is empty")

    columns: dict[int, str] = {}
    extras: dict[int, str] = {}
    for index, cell in enumerate(header):
        name = (_text(cell) or "").lower()
        if name in _HEADERS:
            columns[index] = _HEADERS[name]
        elif name.startswith(EXTRA_PREFIX):
            extras[index] = (_text(cell) or "")[len(EXTRA_PREFIX) :].strip()
    if not {"room", "description"} <= set(columns.values()):
        raise ImportFileError(
            "Columns 'Room' and 'Description' are required. "
            "Use a file exported from Versions."
        )

    parsed: list[dict] = []
    for number, row in enumerate(rows, start=2):
        if not any(_text(c) for c in row):
            continue
        if len(parsed) >= MAX_ROWS:
            raise ImportFileError(f"More than {MAX_ROWS} rows")
        item: dict = {"extra": {}}
        for index, field in columns.items():
            value = row[index] if index < len(row) else None
            if field in _NUMBER_FIELDS:
                item[field] = _number(value, field, number)
            elif field == "status":
                item[field] = _status(value, number)
            else:
                item[field] = _text(value)
        for index, key in extras.items():
            value = row[index] if index < len(row) else None
            if key and _text(value) is not None:
                item["extra"][key] = _text(value)

        for required in ("room", "description"):
            if not item.get(required):
                raise ImportFileError(f"Row {number}: {required} is empty")
        for field, limit in _MAX_LEN.items():
            if item.get(field) and len(item[field]) > limit:
                raise ImportFileError(
                    f"Row {number}: {field.replace('_', ' ')} is longer "
                    f"than {limit} characters"
                )
        item["unit"] = normalize_unit(item.get("unit")) or _DEFAULT_UNIT
        if item.get("status") is None:
            item.pop("status", None)
        parsed.append(item)

    if not parsed:
        raise ImportFileError("The file has no line items")
    return parsed
